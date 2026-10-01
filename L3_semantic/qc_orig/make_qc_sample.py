#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""分层抽取质检样本：盲标表 + 含模型答案的对照表。"""

from __future__ import annotations

import argparse
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

TZ = timezone(timedelta(hours=8))
ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT.parent / "data/raw/taptap"
ANN = ROOT
QC = ROOT / "qc"

SENT_CHOICES = "正,中,负"
# 下拉用中文；打分时由 score_qc 映射为英文码
RHET_CHOICES = "无修辞,讽刺,高级黑,反串,模板好评"
RHETORIC_OK = {"none", "sarcasm", "gaoji_hei", "fanchuan", "template_praise"}


def load_pool(ann_path: Path) -> pd.DataFrame:
    clean = pd.read_csv(ROOT / "data/processed/reviews" / "reviews_clean.csv", dtype={"review_id": str})
    ann = pd.read_csv(ann_path, dtype={"review_id": str})
    ann = ann[ann["error"].fillna("") == ""].drop_duplicates("review_id", keep="last")
    ann = ann[~ann["review_id"].astype(str).str.startswith("probe_")]
    m = ann.merge(
        clean[["review_id", "text", "publish_time_cn"]],
        on="review_id",
        how="left",
    )
    m["score_raw"] = m["score_raw"].astype(str)
    m["rhetoric"] = m["rhetoric"].fillna("none").replace("", "none")
    m["sentiment"] = m["sentiment"].fillna("中")
    m["text"] = m["text"].fillna("").astype(str)
    m = m[m["text"].str.strip() != ""].copy()
    return m


def take(df: pd.DataFrame, n: int, rng) -> pd.DataFrame:
    if len(df) <= n:
        return df.copy()
    return df.sample(n=n, random_state=rng)


def build_sample(pool: pd.DataFrame, n: int, seed: int) -> pd.DataFrame:
    """修辞过采样 + 高星负向 + 情绪分层补齐。"""
    rng = seed
    picked_ids: set[str] = set()
    parts: list[pd.DataFrame] = []

    def add(df: pd.DataFrame, k: int, tag: str) -> None:
        nonlocal rng
        if k <= 0 or df.empty:
            return
        sub = df[~df["review_id"].isin(picked_ids)]
        got = take(sub, k, rng)
        rng += 1
        if got.empty:
            return
        got = got.copy()
        got["strata"] = tag
        parts.append(got)
        picked_ids.update(got["review_id"].tolist())

    rhe = pool[pool["rhetoric"] != "none"]
    # 修辞：目标约 40，尽量覆盖各类
    add(rhe[rhe["rhetoric"] == "sarcasm"], 28, "rhetoric_sarcasm")
    add(rhe[rhe["rhetoric"] == "gaoji_hei"], 8, "rhetoric_gaoji_hei")
    add(rhe[rhe["rhetoric"].isin(["template_praise", "fanchuan"])], 4, "rhetoric_other")
    # 若不足 40，用任意 rhetoric≠none 补
    need_rhe = max(0, 40 - len(picked_ids))
    add(rhe, need_rhe, "rhetoric_fill")

    hi_neg = pool[
        (pool["score_raw"].isin(["4", "5"])) & (pool["sentiment"] == "负")
    ]
    add(hi_neg, 15, "hi_star_neg")

    # 剩余按情绪分层（正:负:中 ≈ 5:4:1）
    remain = n - len(picked_ids)
    if remain > 0:
        n_pos = int(round(remain * 0.5))
        n_neg = int(round(remain * 0.4))
        n_mid = remain - n_pos - n_neg
        base = pool[pool["rhetoric"] == "none"]
        add(base[base["sentiment"] == "正"], n_pos, "sent_pos")
        add(base[base["sentiment"] == "负"], n_neg, "sent_neg")
        add(base[base["sentiment"] == "中"], n_mid, "sent_mid")
        # 仍不足则全局补
        still = n - len(picked_ids)
        add(pool, still, "fill")

    out = pd.concat(parts, ignore_index=True)
    out = out.drop_duplicates("review_id", keep="first")
    if len(out) > n:
        out = out.sample(n=n, random_state=seed).reset_index(drop=True)
    # 打乱展示顺序，避免按 strata 成块暗示
    out = out.sample(frac=1.0, random_state=seed + 99).reset_index(drop=True)
    out.insert(0, "qc_id", range(1, len(out) + 1))
    return out


def write_blind_xlsx(blind: pd.DataFrame, path: Path) -> None:
    """写出带下拉选项的盲标表（Excel）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "盲标"

    headers = [
        "qc_id",
        "review_id",
        "score_raw",
        "text",
        "gold_sentiment",
        "gold_rhetoric",
        "gold_note",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for _, row in blind.iterrows():
        ws.append(
            [
                int(row["qc_id"]),
                str(row["review_id"]),
                str(row["score_raw"]),
                str(row["text"]),
                str(row.get("gold_sentiment") or ""),
                str(row.get("gold_rhetoric") or ""),
                str(row.get("gold_note") or ""),
            ]
        )

    n = len(blind)
    last = n + 1  # header + rows

    dv_sent = DataValidation(
        type="list",
        formula1=f'"{SENT_CHOICES}"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="情绪取值",
        error="请选：正 / 中 / 负",
    )
    dv_rhe = DataValidation(
        type="list",
        formula1=f'"{RHET_CHOICES}"',
        allow_blank=True,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="修辞取值",
        error="请选：无修辞 / 讽刺 / 高级黑 / 反串 / 模板好评",
    )
    ws.add_data_validation(dv_sent)
    ws.add_data_validation(dv_rhe)
    dv_sent.add(f"E2:E{last}")
    dv_rhe.add(f"F2:F{last}")

    # 说明页
    ws2 = wb.create_sheet("填写说明", 0)
    tips = [
        ["列", "怎么填"],
        ["gold_sentiment", "下拉选：正 / 中 / 负（整体态度，不是星级）"],
        ["gold_rhetoric", "下拉选：无修辞 / 讽刺 / 高级黑 / 反串 / 模板好评"],
        ["gold_note", "可选备注"],
        ["", ""],
        ["提示", "请在「盲标」工作表用下拉填写；不要打开 qc_with_model.csv"],
        ["默认", "修辞拿不准 → 无修辞；夸中带槽偏正；五星辱骂 → 负"],
        ["转化", "打分时会自动把中文修辞映射为 none/sarcasm/… 英文码"],
        ["", ""],
        ["中文选项", "对应英文码（打分用，标注不用管）"],
        ["无修辞", "none"],
        ["讽刺", "sarcasm"],
        ["高级黑", "gaoji_hei"],
        ["反串", "fanchuan"],
        ["模板好评", "template_praise"],
    ]
    for r in tips:
        ws2.append(r)
    ws2["A1"].font = Font(bold=True)
    ws2["B1"].font = Font(bold=True)
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 72

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 24
    for row in ws.iter_rows(min_row=2, max_row=last, min_col=4, max_col=4):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.auto_filter.ref = f"A1:G{last}"
    ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> int:
    p = argparse.ArgumentParser(description="Make QC sample for human gold labels")
    p.add_argument("--ann", default=str(ANN / "annotations_v1_4.csv"))
    p.add_argument("--n", type=int, default=100)
    p.add_argument("--seed", type=int, default=42)
    args = p.parse_args()

    pool = load_pool(Path(args.ann))
    sample = build_sample(pool, args.n, args.seed)
    QC.mkdir(parents=True, exist_ok=True)

    # 盲标表：不含模型答案
    blind = sample[
        ["qc_id", "review_id", "score_raw", "text"]
    ].copy()
    blind["gold_sentiment"] = ""
    blind["gold_rhetoric"] = ""
    blind["gold_note"] = ""
    blind_path = QC / "qc_blind_to_label.csv"
    blind.to_csv(blind_path, index=False, encoding="utf-8-sig")
    xlsx_path = QC / "qc_blind_to_label.xlsx"
    write_blind_xlsx(blind, xlsx_path)

    # 对照表：含模型答案（打分用，标注时别打开）
    model = sample[
        [
            "qc_id",
            "review_id",
            "score_raw",
            "text",
            "sentiment",
            "rhetoric",
            "rhetoric_confidence",
            "incongruity_cues",
            "reason",
            "topic_primary",
            "strata",
            "prompt_version",
            "model",
        ]
    ].rename(
        columns={
            "sentiment": "model_sentiment",
            "rhetoric": "model_rhetoric",
        }
    )
    model_path = QC / "qc_with_model.csv"
    model.to_csv(model_path, index=False, encoding="utf-8-sig")

    # 元信息
    meta = QC / "qc_sample_meta.md"
    rhe_n = int((sample["rhetoric"] != "none").sum())
    hi_n = int(
        ((sample["score_raw"].isin(["4", "5"])) & (sample["sentiment"] == "负")).sum()
    )
    meta.write_text(
        "\n".join(
            [
                "# 质检样本元信息",
                "",
                f"- 生成时间：{datetime.now(TZ).isoformat(timespec='seconds')}",
                f"- 标注源：`{Path(args.ann).name}`",
                f"- 样本量：**{len(sample)}**（seed={args.seed}）",
                f"- 其中 rhetoric≠none：**{rhe_n}**",
                f"- 其中高星负向（模型）：**{hi_n}**",
                f"- 情绪分布（模型）：{sample['sentiment'].value_counts().to_dict()}",
                f"- 修辞分布（模型）：{sample['rhetoric'].value_counts().to_dict()}",
                f"- strata：{sample['strata'].value_counts().to_dict()}",
                "",
                "## 文件",
                "",
                "| 文件 | 用途 |",
                "|------|------|",
                "| `qc_blind_to_label.xlsx` | **推荐：带下拉选项，你来填** |",
                "| `qc_blind_to_label.csv` | 同内容纯文本备份 |",
                "| `qc_with_model.csv` | 打分对照（标注时勿看） |",
                "",
                "填完后运行：`python L3_semantic/qc_orig/score_qc.py`",
                "",
            ]
        ),
        encoding="utf-8",
    )

    print(xlsx_path)
    print(blind_path)
    print(model_path)
    print(meta)
    print(f"n={len(sample)} rhe≠none={rhe_n} hi_neg_model={hi_n}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
